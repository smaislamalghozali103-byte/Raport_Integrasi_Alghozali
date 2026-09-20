from io import BytesIO
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

def make_template(students,unit,kelas,mapel,tahun):
    wb=Workbook(); ws=wb.active; ws.title="NILAI"
    ws["A1"]="RAPORT INTEGRASI PONDOK MODERN AL-GHOZALI"
    meta=[("Unit",unit),("Kelas",kelas),("Mata Pelajaran",mapel),("Tahun Ajaran",tahun)]
    for i,(a,b) in enumerate(meta,2): ws.cell(i,1,a); ws.cell(i,2,b)
    headers=["No","ID_SISWA","NIS","NISN","Nama Siswa","Nilai"]
    for c,h in enumerate(headers,1):
        x=ws.cell(7,c,h); x.font=Font(bold=True,color="FFFFFF"); x.fill=PatternFill("solid",fgColor="166534")
        x.alignment=Alignment(horizontal="center")
    for i,s in enumerate(students,1):
        vals=[i,s["id"],s["nis"] or "",s["nisn"] or "",s["nama"],""]
        for c,v in enumerate(vals,1): ws.cell(7+i,c,v)
    widths=[8,14,16,16,38,14]
    for i,w in enumerate(widths,1): ws.column_dimensions[chr(64+i)].width=w
    ws.freeze_panes="A8"
    out=BytesIO(); wb.save(out); return out.getvalue()

def read_uploaded(file):
    df=pd.read_excel(file); df.columns=[str(c).strip() for c in df.columns]
    def col(names):
        return next((x for x in names if x in df.columns),None)
    cval=col(["Nilai","NILAI","Score","score"])
    cid=col(["ID_SISWA","id_siswa","ID"])
    cnisn=col(["NISN","nisn"]); cname=col(["Nama Siswa","Nama","NAMA SISWA","nama"])
    if not cval: raise ValueError("Kolom NILAI tidak ditemukan.")
    out=[]
    for _,r in df.iterrows():
        raw=r[cval]
        if pd.isna(raw) or str(raw).strip()=="": val=None
        else:
            val=float(raw)
            if not 0 <= val <= 100: raise ValueError(f"Nilai {val} di luar rentang 0-100.")
        out.append({"id":r[cid] if cid else None,"nisn":str(r[cnisn]).strip() if cnisn and not pd.isna(r[cnisn]) else "",
                    "nama":str(r[cname]).strip() if cname and not pd.isna(r[cname]) else "","nilai":val})
    return out
